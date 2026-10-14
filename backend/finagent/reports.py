"""Excel exports are rendered from owned immutable artifacts, never chat prose."""

import asyncio
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .contracts import Artifact, DocumentCitation, ExportRequest, FinanceError
from .repository import get_artifact, report_root, save_artifact
from .warnings import artifact_warnings

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def safe_text(value) -> str:
    text = str(value if value is not None else "")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    if text.lstrip().startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text[:32767]


def excel_number(value: str | None):
    if value is None:
        return None
    try:
        number = Decimal(value)
    except InvalidOperation as exc:
        raise FinanceError("invalid_number", "报表输入包含无效数值。") from exc
    if not number.is_finite():
        raise FinanceError("invalid_number", "报表输入包含非有限数值。")
    # Excel stores only 15 significant digits; retain a textual exact value
    # alongside displayable numeric cells rather than silently rounding facts.
    if len(number.normalize().as_tuple().digits) > 15:
        return str(number)
    return float(number)


def _style(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B53")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in sheet.columns:
        width = min(65, max(12, max(len(str(c.value or "")) for c in column) + 2))
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_workbook(path: Path, title: str, data: list[Artifact], analyses: list[Artifact],
                   documents: list[Artifact] | None = None) -> None:
    documents = documents or []
    workbook = Workbook()
    overview = workbook.active
    overview.title = "说明"
    overview.append(["项目", "内容"])
    overview.append(["标题", safe_text(title)])
    overview.append(["依据", "已保存的 SEC 原始申报文档版本及所列结构化数值（如有）；下载或引用不代表已通读全文，文档分析不自动变成已验证的结构化财务事实。" if documents else "官方 SEC XBRL 财务事实；原文链接不代表已下载或阅读原文。"])
    overview.append(["期间", "使用各公司的实际财季起止日期。"])
    overview.append(["精度", "数值来自已保存的结构化结果；精确值列保留原始十进制文本，超过 Excel 15 位精度的值以文本保存。"])
    sources = workbook.create_sheet("来源")
    sources.append(["结果ID", "证据ID", "事实ID", "公司", "XBRL概念", "开始", "结束", "精确值", "单位", "申报", "申报日期", "官方API", "原文链接", "原文状态"])
    coverage = workbook.create_sheet("覆盖与限制")
    coverage.append(["结果ID", "公司", "期间", "指标", "状态", "原因或限制"])
    all_warnings = artifact_warnings(data + analyses + documents)
    if documents:
        document_sources = workbook.create_sheet("文档来源")
        document_sources.append(["文档结果ID", "来源文档ID", "公司", "CIK", "申报类型", "报告期末", "Accession", "主文档", "来源URL", "SHA-256", "字节数", "获取时间", "申报日期", "状态"])
        for artifact in documents:
            document = artifact.document
            if document is None:
                continue
            document_sources.append([
                artifact.artifact_id, document.source_document_id, document.company, document.cik,
                document.form, str(document.period_end or ""), document.accession, document.primary_document,
                document.source_url, document.sha256, document.byte_size, document.fetched_at.isoformat(), str(document.filed_at or ""),
                "已下载固定版本；不代表已通读全文",
            ])
            overview.append([artifact.artifact_id, safe_text(artifact.title)])
    for index, artifact in enumerate(data, 1):
        table = artifact.data
        if table is None:
            continue
        sheet = workbook.create_sheet(f"数据{index}")
        sheet.append(["结果ID", "公司", "财季", "开始", "结束", "指标", "值", "精确值", "单位", "状态", "公式", "证据ID", "单元格ID"])
        requested = table.specification.get("requested_metrics") or []
        for row in table.rows:
            for cell in row.cells:
                if requested and cell.metric not in requested:
                    continue
                sheet.append([
                    artifact.artifact_id, safe_text(row.company), safe_text(row.period_label),
                    str(row.period_start or ""), str(row.period_end), safe_text(cell.label),
                    excel_number(cell.value), safe_text(cell.value), safe_text(cell.unit), cell.status,
                    safe_text(cell.formula), safe_text(", ".join(cell.source_ids)), safe_text(cell.cell_id),
                ])
                coverage.append([artifact.artifact_id, row.company, str(row.period_end), safe_text(cell.label), cell.status, safe_text(cell.reason)])
        for evidence in table.sources.values():
            sources.append([
                artifact.artifact_id, evidence.source_id, evidence.fact_id, evidence.company,
                safe_text(evidence.concept), str(evidence.period_start or ""), str(evidence.period_end),
                safe_text(evidence.value), safe_text(evidence.unit), safe_text(evidence.accession),
                str(evidence.filed_at), safe_text(evidence.api_url), safe_text(evidence.filing_url), evidence.document_status,
            ])
        overview.append([artifact.artifact_id, safe_text(artifact.title)])
    for index, artifact in enumerate(analyses, 1):
        sheet = workbook.create_sheet(f"分析{index}")
        sheet.append(["项目", "内容"])
        sheet.append(["结果ID", artifact.artifact_id])
        sheet.append(["标题", safe_text(artifact.title)])
        for line in (artifact.text or "").splitlines():
            sheet.append(["正文", safe_text(line)])
        for citation in artifact.citations:
            if isinstance(citation, DocumentCitation):
                sheet.append(["文档引用", safe_text(f"{citation.document_id} / {citation.label}")])
            else:
                sheet.append(["引用", safe_text(f"{citation.artifact_id} / {citation.cell_id} / {citation.label}")])
    for warning in all_warnings:
        coverage.append(["", "", "", "", "warning", safe_text(warning)])
    for sheet in workbook.worksheets:
        # Sanitize every textual field, including identifiers and scope labels,
        # not just prose. Never leave a user-controlled Excel formula behind.
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = safe_text(cell.value)
        _style(sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)


async def generate_report(user_id: str, run_id: str, request: ExportRequest) -> Artifact:
    data = {}
    documents = {}
    analyses = []
    for artifact_id in request.data_ids:
        artifact = await get_artifact(user_id, artifact_id)
        if artifact.data is None or artifact.kind not in {"dataset", "table"}:
            raise FinanceError("invalid_result", "data_ids 必须引用结构化数据或计算结果。")
        data[artifact_id] = artifact
    for artifact_id in request.analysis_ids:
        artifact = await get_artifact(user_id, artifact_id)
        if artifact.kind != "narrative":
            raise FinanceError("invalid_result", "analysis_ids 必须引用已保存的分析结果。")
        analyses.append(artifact)
        # Resolve both citation variants and every declared dependency through the
        # owned repository. Never infer a document citation's nonexistent cell_id.
        citation_ids = [c.document_id if isinstance(c, DocumentCitation) else c.artifact_id for c in artifact.citations]
        for source_id in dict.fromkeys(artifact.input_ids + citation_ids):
            source = await get_artifact(user_id, source_id)
            if source.data and source.kind in {"dataset", "table"}:
                data[source_id] = source
            elif source.kind == "document" and source.document:
                documents[source_id] = source
            else:
                raise FinanceError("invalid_result", "分析依赖必须为已保存的数值结果或原始申报文档。")
        for citation in artifact.citations:
            if isinstance(citation, DocumentCitation):
                valid = citation.document_id in documents
            else:
                source = data.get(citation.artifact_id)
                valid = source is not None and any(
                    cell.cell_id == citation.cell_id for row in source.data.rows + source.data.history_rows for cell in row.cells
                )
            if not valid:
                raise FinanceError("invalid_citation", "分析引用与所拥有的来源结果不匹配。")
    if not data and not documents:
        raise FinanceError("no_data", "没有可导出的结构化数据或原始申报文档依据。")
    cells = sum(len(row.cells) for artifact in data.values() for row in artifact.data.rows)
    if len(data) + len(documents) > 24 or cells > 10000:
        raise FinanceError("report_too_large", "本次导出范围过大，请缩小结果集合。")
    filename = "finance-report-" + uuid4().hex[:12] + ".xlsx"
    path = report_root() / filename
    await asyncio.to_thread(write_workbook, path, request.title, list(data.values()), analyses, list(documents.values()))
    # A cancelled task never reaches publication. A thread may finish its local
    # file, but no owned download artifact is announced for that orphan file.
    warnings = artifact_warnings(list(data.values()) + analyses + list(documents.values()))
    return await save_artifact(
        user_id, run_id, "export", request.title, warnings=warnings,
        input_ids=list(dict.fromkeys(list(data) + list(documents) + request.analysis_ids)), filename=filename,
        mime_type=MIME_XLSX, path=str(path),
    )
