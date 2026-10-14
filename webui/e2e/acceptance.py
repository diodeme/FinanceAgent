"""Opt-in browser acceptance; defaults to inspection and never provisions credentials.

Build: docker build -t finance-e2e webui/e2e
Run with FIN_AGENT_ID/FIN_SESSION_ID and mount ignored host data/e2e-output at
/e2e-output. Container URLs must be reachable from the browser container (for
example http://host.docker.internal:8000 and :8080 on Docker Desktop).

Real model usage requires --multi-turn. Override the four default prompts with
repeated --prompt, --prompts-file (JSON string array), or FIN_PROMPTS_JSON.
There are no automatic chat retries, background model calls, or session creation.
"""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import re
import sys
import time
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from urllib.parse import quote, urlsplit

import httpx
from openpyxl import load_workbook
from playwright.sync_api import Page, sync_playwright


DEFAULT_PROMPTS = [
    "查一下苹果最近一个已披露财季的营收，显示实际起止日期、单位和可点击的来源。",
    "再对比苹果和微软最近四个已披露财季的营收和毛利率，保留各公司实际财季日期。",
    "基于刚才的表格做一段简短分析，引用具体数据；不推测没有原文证据支持的原因。",
    "把刚才的对比表格和分析整理成一个 Excel 报告供我下载，复用已有结果。",
]
ARTIFACT = "[data-finance-artifact]"
ASSISTANT = '[data-slot="message"][data-role="assistant"]'
STOP = 'button[data-size="icon-lg"]:has(svg.lucide-square)'
DRAWER = '[data-testid="finance-provenance"]'


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def sanitized_error(error: object) -> str:
    # Never persist headers, response bodies, Playwright call logs, or keys.
    text = str(error).split("\n")[0]
    text = re.sub(r"(?i)(bearer\s+|(?:api[_-]?key|authorization|token)[=: ]+)[^\s,;]+", r"\1[REDACTED]", text)
    text = re.sub(r"\bsk-[A-Za-z0-9_-]+", "[REDACTED]", text)
    text = re.sub(r"(https?://[^\s?#]+)[?#][^\s]+", r"\1?[REDACTED]", text)
    return text[:400]


def base_url(value: str) -> str:
    parsed = urlsplit(value)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password or parsed.query or parsed.fragment:
        raise argparse.ArgumentTypeError("Use an HTTP(S) application URL without embedded credentials, query or fragment")
    return value.rstrip("/")


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--inspect-only", action="store_true", help="Inspect existing UI/results only; never send a chat (default)")
    mode.add_argument("--multi-turn", action="store_true", help="Explicitly authorize this test to send the configured prompts")
    parser.add_argument("--api-url", type=base_url, default=os.getenv("FIN_API_URL", "http://127.0.0.1:8000"))
    parser.add_argument("--web-url", type=base_url, default=os.getenv("FIN_WEB_URL", "http://127.0.0.1:8080"))
    parser.add_argument("--user-id", default=os.getenv("FIN_USER_ID", "dev"), help="Explicit existing test user; default FIN_USER_ID or dev")
    parser.add_argument("--agent-id", default=os.getenv("FIN_AGENT_ID"))
    parser.add_argument("--session-id", default=os.getenv("FIN_SESSION_ID"))
    parser.add_argument("--prompt", action="append", help="Repeat for each turn; no prompt content is written to the report")
    parser.add_argument("--prompts-file", type=Path, help="UTF-8 JSON array of prompts")
    parser.add_argument("--turn-timeout", type=int, default=360, help="Maximum seconds per real turn; no retries")
    parser.add_argument("--ui-timeout", type=int, default=45, help="Maximum seconds for UI navigation/provenance/download")
    parser.add_argument("--output-dir", type=Path, default=Path(os.getenv("FIN_E2E_OUTPUT", "/e2e-output")))
    parser.add_argument("--headed", action="store_true")
    args = parser.parse_args()
    if not args.agent_id or not args.session_id or not args.user_id:
        parser.error("An existing FIN_AGENT_ID, FIN_SESSION_ID and nonempty FIN_USER_ID are required")
    if args.turn_timeout <= 0 or args.ui_timeout <= 0:
        parser.error("Timeouts must be positive")
    if args.prompt and args.prompts_file:
        parser.error("Choose --prompt or --prompts-file, not both")
    try:
        args.prompts = args.prompt or (json.loads(args.prompts_file.read_text(encoding="utf-8")) if args.prompts_file else json.loads(os.getenv("FIN_PROMPTS_JSON", "null"))) or DEFAULT_PROMPTS
    except (ValueError, OSError) as error:
        parser.error(f"Cannot read configured prompts: {type(error).__name__}")
    if not isinstance(args.prompts, list) or not 1 <= len(args.prompts) <= 4 or not all(isinstance(item, str) and item.strip() for item in args.prompts):
        parser.error("Configure between one and four nonempty prompt strings")
    return args


class Acceptance:
    def __init__(self, args: argparse.Namespace):
        self.args = args
        self.output = args.output_dir / datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
        self.output.mkdir(parents=True, exist_ok=False)
        self.report = {
            "mode": "multi-turn" if args.multi_turn else "inspect-only",
            "user_id": args.user_id, "agent_id": args.agent_id, "session_id": args.session_id,
            "status": "running", "turns": [], "checks": {}, "screenshots": [], "page_errors": [],
        }
        self.api = httpx.Client(base_url=args.api_url, headers={"X-User-ID": args.user_id}, timeout=30)
        self.page: Page | None = None

    def screenshot(self, name: str) -> None:
        path = self.output / f"{name}.png"
        self.page.screenshot(path=str(path), full_page=True, timeout=self.args.ui_timeout * 1000)
        self.report["screenshots"].append(str(path))

    def api_get(self, path: str, params: dict | None = None) -> dict:
        # Only fixed, owned application resources; no model-provided URL is fetched.
        response = self.api.get(path, params=params)
        require(response.is_success, f"Owned finance API request failed (HTTP {response.status_code})")
        return response.json()

    def artifacts(self) -> list[str]:
        return self.page.locator(ARTIFACT).evaluate_all("nodes => [...new Set(nodes.map(n => n.dataset.financeArtifact))]")

    def snapshot(self, since: int = 0) -> dict:
        return self.page.evaluate("""({assistant, stop, since}) => {
            const messages = [...document.querySelectorAll(assistant)];
            // A native turn can append several assistant messages. Observe all
            // messages from this turn, never just the last or historical tools.
            const current = messages.slice(since);
            const tools = current.flatMap(message => [...message.querySelectorAll('[data-finance-tool]')]).map(n => ({
                name: n.dataset.financeTool, status: n.dataset.financeStatus,
                textStreaming: ['running', 'saving'].includes(n.dataset.financeStatus) && (
                    [...n.querySelectorAll('[role="log"], .cn-markdown')].some(t => t.textContent.trim().length > 0) ||
                    [...n.querySelectorAll('[data-finance-report-text-length]')].some(t => Number(t.dataset.financeReportTextLength) > 0) ||
                    [...n.querySelectorAll('[data-finance-report-progress-count]')].some(t => Number(t.dataset.financeReportProgressCount) > 0))
            }));
            return {assistantCount: messages.length, stop: !!document.querySelector(stop), tools,
                footerRunning: current.some(message => !!message.querySelector('[data-slot="message-footer"] [data-slot="badge"][aria-label]')),
                nativeError: current.some(message => !!message.querySelector('[role="alert"]'))};
        }""", {"assistant": ASSISTANT, "stop": STOP, "since": since})

    def send_turn(self, index: int, prompt: str) -> None:
        page = self.page
        expected_path = f"/chat/{quote(self.args.agent_id, safe='')}/{quote(self.args.session_id, safe='')}"
        require(urlsplit(page.url).path.rstrip("/") == expected_path, "Browser left the explicitly supplied test session; refusing to send")
        require(not page.locator(STOP).count(), "Session already has an active run; refusing to overlap requests")
        before = page.locator(ASSISTANT).count()
        turn = {"index": index, "prompt_characters": len(prompt), "tools": [], "stream_observed": False, "status": "running"}
        self.report["turns"].append(turn)
        composer = page.locator("textarea:visible").last
        composer.fill(prompt)
        composer.press("Enter")  # Real composer path, not a synthetic /chat API call.
        started = time.monotonic()
        deadline = started + self.args.turn_timeout
        settled_at = None
        seen_tools = set()
        while time.monotonic() < deadline:
            snapshot = self.snapshot(before)
            if snapshot["assistantCount"] > before:
                for tool in snapshot["tools"]:
                    seen_tools.add(tool["name"])
                if not turn["stream_observed"] and any(tool["textStreaming"] for tool in snapshot["tools"]):
                    turn["stream_observed"] = True
                    self.screenshot(f"turn-{index}-streaming")
                finished = not snapshot["stop"] and not snapshot["footerRunning"] and not any(tool["status"] in {"running", "saving"} for tool in snapshot["tools"])
                if finished:
                    settled_at = settled_at or time.monotonic()
                    if time.monotonic() - settled_at >= 0.8:
                        turn.update({"tools": sorted(seen_tools), "duration_seconds": round(time.monotonic() - started, 2), "tool_states": snapshot["tools"]})
                        self.screenshot(f"turn-{index}-finished")
                        failed = snapshot["nativeError"] or any(tool["status"] in {"failed", "interrupted", "incomplete"} for tool in snapshot["tools"])
                        turn["status"] = "failed" if failed else "complete"
                        require(not failed, f"Turn {index} ended with an error/incomplete tool result; no next prompt was sent")
                        return
                else:
                    settled_at = None
            page.wait_for_timeout(200)
        turn.update({"status": "timeout", "tools": sorted(seen_tools), "duration_seconds": self.args.turn_timeout})
        # Stop the one run we started; never retry it or create a replacement.
        stop = page.locator(STOP)
        if stop.count() and stop.last.is_enabled():
            stop.last.click(timeout=5000)
            turn["stop_requested"] = True
        self.screenshot(f"turn-{index}-timeout")
        raise TimeoutError(f"Turn {index} exceeded its bounded timeout")

    def open_source(self, trigger, name: str, *, capture: bool = True) -> dict:
        trigger.scroll_into_view_if_needed()
        trigger.click()
        drawer = self.page.locator(DRAWER)
        drawer.wait_for(state="visible")
        self.page.wait_for_function("""selector => {
            const drawer = document.querySelector(selector);
            return drawer && !drawer.querySelector('[role="status"]') &&
                (drawer.querySelector('dl') || drawer.querySelector('[role="alert"]'));
        }""", arg=DRAWER, timeout=self.args.ui_timeout * 1000)
        require(drawer.locator('[role="alert"]').count() == 0, "Provenance request failed")
        source_count = drawer.locator("dl").count()
        require(source_count > 0, "Provenance drawer has no raw financial facts")
        text = drawer.inner_text()
        statuses = []
        for status, label in {"metadata_only": "仅有元数据和链接，未获取正文", "unavailable": "正文不可用", "fetched": "已获取正文"}.items():
            if label in text:
                statuses.append(status)
        require(bool(statuses), "Original filing availability is not explicit")
        require(drawer.get_by_role("link", name="官方 API 证据").count() > 0, "Official API evidence link missing")
        if capture:
            self.screenshot(name)
        result = {"status": "passed", "source_count": source_count, "document_statuses": statuses, "formula_visible": "计算公式" in text, "inputs_visible": "计算输入" in text}
        self.page.keyboard.press("Escape")
        drawer.wait_for(state="hidden")
        return result

    def test_sources(self, suffix: str = "") -> None:
        buttons = self.page.locator(f"{ARTIFACT} table tbody td button")
        available = None
        for index in range(buttons.count()):
            if not buttons.nth(index).inner_text().lstrip().startswith("—"):
                available = buttons.nth(index)
                break
        if available is not None:
            self.report["checks"][f"table_source{suffix}"] = self.open_source(available, f"table-source{suffix}")
        else:
            self.report["checks"][f"table_source{suffix}"] = {"status": "missing"}
            require(not self.args.multi_turn, "No clickable financial value was rendered")
        report_panel = self.page.locator('[data-testid="finance-report"]')
        report_buttons = self.page.locator(f'{ARTIFACT} [data-finance-report-open]')
        if report_buttons.count():
            require(self.page.locator('[data-finance-report-card] .cn-markdown').count() == 0,
                    "Full narrative is still rendered inside the compact chat card")
            report_buttons.first.click()
            report_panel.wait_for(state="visible")
            links = report_panel.locator('.cn-markdown a[data-finance-citation][href^="#finance/"]')
        else:
            # Older deployments retain the inline report; useful for before/after diagnostics.
            links = self.page.locator(f'{ARTIFACT} .cn-markdown a[data-finance-citation][href^="#finance/"]')
        if links.count():
            self.report["checks"][f"narrative_citation{suffix}"] = self.open_source(links.first, f"narrative-citation{suffix}")
        else:
            self.report["checks"][f"narrative_citation{suffix}"] = {"status": "missing"}
            require(not self.args.multi_turn, "No narrative Markdown citation was rendered")
        if report_buttons.count():
            require(report_panel.is_visible(), "Closing provenance also closed the report")
            self.page.keyboard.press("Escape")
            report_panel.wait_for(state="hidden")
            self.report["checks"][f"document_card{suffix}"] = {"status": "passed", "nested_provenance": True}

    def test_assistant_citations(self, suffix: str = "") -> None:
        history = self.api_get(f"/sessions/{quote(self.args.session_id, safe='')}/messages",
                               {"agent_id": self.args.agent_id, "limit": 100})
        expected = set()
        for message in history.get("messages", []):
            if message.get("role") != "assistant":
                continue
            for block in message.get("content", []):
                if block.get("type") != "text":
                    continue
                for line in block.get("text", "").splitlines():
                    if line.lstrip().startswith("|"):
                        expected.update(re.findall(r"\[([^\[\]\n]+)\]\((#finance/[^\s()]+)\)", line))
        key = f"assistant_table_citations{suffix}"
        if not expected:
            self.report["checks"][key] = {"status": "not_applicable"}
            return
        selector = f'{ASSISTANT} .cn-markdown table a[data-finance-citation]'
        self.page.wait_for_function("({selector, expected}) => { const links = [...document.querySelectorAll(selector)].filter(n => !n.closest('[data-finance-artifact]')); return expected.every(([label, href]) => links.some(n => n.textContent.trim() === label && n.getAttribute('href') === href)); }", arg={"selector": selector, "expected": sorted(expected)})
        anchors = self.page.locator(selector)
        observed = anchors.evaluate_all("nodes => nodes.filter(n => !n.closest('[data-finance-artifact]')).map(n => [n.textContent.trim(), n.getAttribute('href')])")
        require(expected.issubset({tuple(item) for item in observed}),
                "Some persisted assistant-table citations have no interactive rendered anchor")
        checked = []
        for index, (label, href) in enumerate(sorted(expected)):
            parts = href.removeprefix("#finance/").split("/")
            require(len(parts) == 2, "Malformed persisted finance reference")
            path = f"/fin/results/{parts[0]}/cells/{parts[1]}"
            trigger = self.page.locator(f'{ASSISTANT} .cn-markdown table a[data-finance-citation][href={json.dumps(href)}]').first
            with self.page.expect_response(lambda response: urlsplit(response.url).path == path) as pending:
                result = self.open_source(trigger, f"assistant-citation-{index}{suffix}", capture=index < 3)
            response = pending.value
            require(response.status == 200, "Assistant-table citation did not fetch authorized provenance")
            provenance = response.json()
            require(provenance.get("cell", {}).get("cell_id") == parts[1], "Citation opened a different cell")
            checked.append({"label": label, "cell_id": parts[1], **result})
        self.report["checks"][key] = {"status": "passed", "count": len(checked), "citations": checked}

    def test_formula(self) -> None:
        results = self.api_get("/fin/results", {"session_id": self.args.session_id, "limit": "10"})
        for artifact in results.get("artifacts", []):
            data = artifact.get("data") or {}
            rows = data.get("rows", [])
            requested = data.get("specification", {}).get("requested_metrics")
            columns = list(dict.fromkeys(cell["metric"] for row in rows for cell in row["cells"] if not requested or cell["metric"] in requested))
            cards = self.page.locator(ARTIFACT)
            card = next((cards.nth(i) for i in range(cards.count()) if cards.nth(i).get_attribute("data-finance-artifact") == artifact["artifact_id"]), None)
            if card is None:
                continue
            for row_index, row in enumerate(rows):
                for cell in row["cells"]:
                    if not cell.get("formula") or cell["metric"] not in columns or cell.get("value") is None:
                        continue
                    button = card.locator("table tbody tr").nth(row_index).locator("td").nth(columns.index(cell["metric"]) + 2).get_by_role("button")
                    result = self.open_source(button, "formula-inputs")
                    require(result["formula_visible"], "Calculated cell has no visible formula in its drawer")
                    require(not cell.get("inputs") or result["inputs_visible"], "Calculated cell input facts are not visible")
                    self.report["checks"]["formula_inputs"] = result
                    return
        self.report["checks"]["formula_inputs"] = {"status": "not_present"}

    def verify_workbook(self, path: Path, export_id: str) -> dict:
        exported = self.api_get(f"/fin/results/{quote(export_id, safe='')}")
        expected = {}
        for input_id in exported.get("input_ids", []):
            artifact = self.api_get(f"/fin/results/{quote(input_id, safe='')}")
            data = artifact.get("data") or {}
            for row in data.get("rows", []):
                for cell in row.get("cells", []):
                    expected[(input_id, cell["cell_id"])] = cell.get("value")
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            require("来源" in workbook.sheetnames and "覆盖与限制" in workbook.sheetnames, "Workbook lacks evidence/coverage worksheets")
            numeric_cells = 0
            checked_values = 0
            for sheet in workbook:
                iterator = sheet.iter_rows(values_only=True)
                header = next(iterator, ())
                if not {"结果ID", "单元格ID", "值", "精确值"}.issubset(header):
                    continue
                positions = {label: index for index, label in enumerate(header)}
                for row in iterator:
                    raw = row[positions["精确值"]]
                    value = row[positions["值"]]
                    if raw is None:
                        continue
                    key = (row[positions["结果ID"]], row[positions["单元格ID"]])
                    require(key in expected, "Workbook references a cell absent from saved input artifacts")
                    try:
                        # Export text is formula-injection hardened; a leading
                        # apostrophe is an Excel text marker, not a changed sign.
                        exact = Decimal(str(raw).removeprefix("'"))
                        displayed = Decimal(str(value).removeprefix("'"))
                        require(exact == Decimal(expected[key]), "Workbook exact-value column differs from persisted artifact")
                        require(displayed == exact, "Workbook display value lost financial precision")
                    except (InvalidOperation, TypeError):
                        raise AssertionError("Workbook financial value is not a valid decimal") from None
                    checked_values += 1
                    numeric_cells += isinstance(value, (int, float)) and not isinstance(value, bool)
            require(checked_values > 0, "Workbook contains no checked financial values")
            return {"status": "passed", "worksheets": workbook.sheetnames, "checked_values": checked_values, "numeric_cells": numeric_cells, "bytes": path.stat().st_size, "path": str(path)}
        finally:
            workbook.close()

    def test_download(self) -> None:
        buttons = self.page.locator(ARTIFACT).get_by_role("button", name=re.compile(r"下载 Excel|Download Excel"))
        if not buttons.count():
            self.report["checks"]["download"] = {"status": "missing"}
            require(not self.args.multi_turn, "No visible Excel download button")
            return
        button = buttons.last
        export_id = button.evaluate("node => node.closest('[data-finance-artifact]').dataset.financeArtifact")
        button.scroll_into_view_if_needed()
        with self.page.expect_download(timeout=self.args.ui_timeout * 1000) as pending:
            button.click()
        download = pending.value
        require(not download.failure(), "Browser download failed")
        name = Path(download.suggested_filename.replace("\\", "/")).name
        require(name.lower().endswith(".xlsx"), "Download filename is not .xlsx")
        path = self.output / name
        download.save_as(str(path))
        require(path.stat().st_size > 0, "Downloaded workbook is empty")
        self.report["checks"]["download"] = self.verify_workbook(path, export_id)
        self.screenshot("download-complete")

    def run(self) -> int:
        playwright = None
        browser = None
        try:
            initial_results = self.api_get("/fin/results", {"session_id": self.args.session_id, "limit": "10"})
            initial_ids = [artifact["artifact_id"] for artifact in initial_results.get("artifacts", [])]
            playwright = sync_playwright().start()
            browser = playwright.chromium.launch(headless=not self.args.headed)
            context = browser.new_context(locale="zh-CN", viewport={"width": 1440, "height": 1000}, accept_downloads=True)
            storage = json.dumps({"server_url": self.args.api_url, "username": self.args.user_id, "i18nextLng": "zh"})
            context.add_init_script(f"for (const [key, value] of Object.entries({storage})) localStorage.setItem(key, value);")
            self.page = context.new_page()
            self.page.set_default_timeout(self.args.ui_timeout * 1000)
            self.page.on("pageerror", lambda error: self.report["page_errors"].append(sanitized_error(error)))
            route = f"{self.args.web_url}/chat/{quote(self.args.agent_id, safe='')}/{quote(self.args.session_id, safe='')}"
            self.page.goto(route, wait_until="domcontentloaded")
            self.page.locator("textarea:visible").last.wait_for()
            if initial_ids:
                # Composer and agent/session hydration can settle at different times.
                self.page.wait_for_function("ids => { const present = new Set([...document.querySelectorAll('[data-finance-artifact]')].map(n => n.dataset.financeArtifact)); return ids.every(id => present.has(id)); }", arg=initial_ids)
            self.screenshot("initial")
            if self.args.multi_turn:
                for index, prompt in enumerate(self.args.prompts, 1):
                    self.send_turn(index, prompt)
            # Inspect mode never sends, interrupts, switches model, or changes credentials.
            self.report["artifact_ids"] = self.artifacts()
            self.report["tool_names"] = sorted(set(self.page.locator("[data-finance-tool]").evaluate_all("nodes => nodes.map(n => n.dataset.financeTool)")))
            self.test_sources()
            self.test_assistant_citations()
            self.test_formula()
            self.test_download()
            before = self.artifacts()
            self.page.reload(wait_until="domcontentloaded")
            self.page.locator("textarea:visible").last.wait_for()
            if before:
                self.page.wait_for_function("ids => { const present = new Set([...document.querySelectorAll('[data-finance-artifact]')].map(n => n.dataset.financeArtifact)); return ids.every(id => present.has(id)); }", arg=before)
            self.screenshot("reloaded")
            self.report["checks"]["reload"] = {"status": "passed", "restored_artifact_ids": self.artifacts()}
            if before:
                self.test_sources("-after-reload")
                self.test_assistant_citations("-after-reload")
            if self.args.multi_turn:
                observed = any(turn["stream_observed"] for turn in self.report["turns"])
                self.report["checks"]["stream_before_completion"] = {"status": "passed" if observed else "failed"}
                require(observed, "No tool progress/analysis text was observed before native completion")
            require(not self.report["page_errors"], "Browser emitted uncaught page errors")
            self.report["status"] = "passed"
        except Exception as error:
            self.report["status"] = "failed"
            self.report["failure"] = {"type": type(error).__name__, "message": sanitized_error(error)}
            if self.page and not self.page.is_closed():
                try:
                    self.screenshot("failure")
                except Exception:
                    pass
        finally:
            # Preserve the sanitized report even after a Chromium/process crash.
            for resource, close in ((browser, "close"), (playwright, "stop"), (self.api, "close")):
                if resource:
                    try:
                        getattr(resource, close)()
                    except Exception:
                        pass
            report_path = self.output / "report.json"
            report_path.write_text(json.dumps(self.report, ensure_ascii=False, indent=2), encoding="utf-8")
            print(json.dumps({"status": self.report["status"], "report": str(report_path), "mode": self.report["mode"], "user_id": self.args.user_id}, ensure_ascii=False))
        return 0 if self.report["status"] == "passed" else 1


if __name__ == "__main__":
    sys.exit(Acceptance(arguments()).run())
