"""Opt-in, ONE browser prompt for the financial-analysis Skill's full result chain.

No live action occurs on import or without --run-real-model. Inherits the current
model/credential *reference* from an existing session, creates exactly one new
session, and never retries a paid turn. max_tokens is a per-call ceiling, NOT a
whole-turn spending cap. Native model-internal retries are outside this runner.

Reuses acceptance.py for browser lifecycle, numeric citations, formulas, Excel
precision and reload; filing_acceptance.py for original-document download/hash.
A passing report covers the browser/result chain only. Skill invocation/content
hash and report quality require the separately recorded parent audit; merely
finding a workspace Skill does not prove that the model read or followed it.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from contextlib import contextmanager
import json
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook

from acceptance import Acceptance, base_url, require, sanitized_error
from filing_acceptance import FilingAcceptance


SESSION_NAME = "Skill 综合财务分析实测"
MAX_TOKENS = 4000
BASE_METRICS = {
    "revenue", "cost_of_revenue", "operating_income", "net_income",
    "operating_cash_flow", "capital_expenditures", "total_assets",
    "total_liabilities", "total_equity", "cash_and_equivalents",
}
DERIVED_METRICS = {"gross_margin", "operating_margin", "net_margin", "free_cash_flow"}
PROMPT = """请对苹果（AAPL）最新一个已披露财季做综合财务分析，并展示最近四个已披露财季的趋势。如果有 financial-analysis Skill，请先阅读并按它的分析方法完成任务。
请通过财务数据工具取得十项基础指标：营收、营业成本、营业利润、净利润、经营现金流、资本开支、总资产、总负债、股东权益及现金与现金等价物；基础指标一次不要超过十二项。毛利率、营业利润率、净利率、自由现金流以及营收增长请交给确定性计算工具，增长优先同比，无可比基期则明确缺口，不要心算补数。保留实际财季起止日、单位和现金流口径，资产负债表使用期末时点；累计现金流不得直接冒充单季。
同时获取与最新分析季度对应的一份官方原始申报文件（10-Q 或 10-K），不要为了四季趋势额外下载其他原文。结合管理层讨论和风险披露解释数字，明确哪些是原文陈述、哪些是分析推断。缺失数据、口径差异或证据不足要直接说明，不要虚构原因。
请用分析工具把已保存的数值结果和原文文件共同作为依据，形成约600至900字中文报告，覆盖增长、盈利、现金流、资产负债状况和主要风险，报告中同时提供可点击的数值来源与原文引用。最后复用这些已保存的表格、分析和原文来源，导出可下载的 Excel 报告，保留精确数值及文档来源信息。聊天里简短提示完成情况即可，不必再复制整篇报告。如果工具或分析失败，请如实报告并停止，不要自动重复付费分析，也不要另开任务补跑。"""


def model_reference(source: dict) -> dict:
    """Copy only native model reference fields, without accessing credentials."""
    require(not source.get("is_running"), "Source session is running; refusing to overlap")
    config = source.get("session", {}).get("config", {}).get("chat_model_config")
    require(isinstance(config, dict), "Source session has no configured model")
    require(all(isinstance(config.get(key), str) and config[key] for key in
                ("type", "credential_id", "model")), "Source model reference is incomplete")
    parameters = deepcopy(config.get("parameters") or {})
    require(isinstance(parameters, dict), "Model parameters must be an object")
    # Credentials belong in credential_id, not in copied provider parameters.
    def no_secrets(value):
        if isinstance(value, dict):
            require(not any(str(key).lower() in {
                "api_key", "apikey", "authorization", "access_token", "secret", "password",
            } for key in value), "Model parameters contain inline credentials; refusing to copy")
            for child in value.values():
                no_secrets(child)
        elif isinstance(value, list):
            for child in value:
                no_secrets(child)
    no_secrets(parameters)
    current = parameters.get("max_tokens")
    parameters["max_tokens"] = min(current, MAX_TOKENS) if type(current) is int and current > 0 else MAX_TOKENS
    return {**{key: config[key] for key in ("type", "credential_id", "model")}, "parameters": parameters}


def validate_chain(artifacts: list[dict]) -> dict:
    """Owned API artifacts, not chat claims, establish the saved dependency chain."""
    groups = {kind: [item for item in artifacts if item.get("kind") == kind]
              for kind in ("dataset", "table", "document", "narrative", "export")}
    numeric = groups["dataset"] + groups["table"]
    require(groups["dataset"] and groups["table"], "Missing fetched dataset or deterministic table")
    require(len(groups["document"]) == 1, "Expected exactly one original filing; no extra downloads")
    require(len(groups["narrative"]) == 1, "Expected exactly one saved analysis; no analysis retry")
    require(len(groups["export"]) == 1, "Expected exactly one saved Excel export")
    numeric_ids = {item["artifact_id"] for item in numeric}
    document, narrative, export = (groups[kind][0] for kind in ("document", "narrative", "export"))
    inputs = set(narrative.get("input_ids", []))
    require(bool(inputs & numeric_ids) and document["artifact_id"] in inputs,
            "Narrative input_ids must include both numeric data and the original document")
    require(bool(narrative.get("text", "").strip()), "Saved analysis has no report text")
    citations = narrative.get("citations", [])
    require(any(item.get("artifact_id") in numeric_ids and item.get("cell_id") for item in citations),
            "Saved report lacks a numeric citation")
    require(any(item.get("document_id") == document["artifact_id"] for item in citations),
            "Saved report lacks its original-document citation")
    exported = set(export.get("input_ids", []))
    require(bool(numeric_ids & exported) and (inputs & numeric_ids).issubset(exported)
            and {document["artifact_id"], narrative["artifact_id"]}.issubset(exported),
            "Excel must include the analysis's numeric inputs, original document and saved analysis")
    require(export.get("filename", "").lower().endswith(".xlsx"), "Export is not an Excel workbook")

    rows = [row for item in groups["dataset"] for row in (item.get("data") or {}).get("rows", [])]
    require(rows and all(row.get("company") == "AAPL" for row in rows), "Expected AAPL numeric rows only")
    periods = sorted({(row.get("period_start"), row.get("period_end")) for row in rows})
    require(len(periods) == 4 and all(start and end and start < end for start, end in periods),
            "Fetched data must retain four distinct actual fiscal-quarter start/end dates")
    for period in periods:
        metrics = {cell["metric"] for row in rows if (row["period_start"], row["period_end"]) == period
                   for cell in row.get("cells", [])}
        require(BASE_METRICS.issubset(metrics), "A fiscal quarter is missing requested base-metric entries")
    require(document.get("document", {}).get("company") == "AAPL", "Original document is not for AAPL")
    require(document["document"].get("period_end") == max(end for _, end in periods),
            "Original filing and latest numeric quarter are not aligned")
    calculated = [item for item in groups["table"]
                  if (item.get("data") or {}).get("specification", {}).get("operation") == "calculate"]
    cells = [cell for item in calculated for row in item["data"].get("rows", []) for cell in row.get("cells", [])]
    available = {cell["metric"] for cell in cells if cell.get("formula") and cell.get("inputs")
                 and cell.get("value") is not None}
    require(DERIVED_METRICS.issubset(available), "Deterministic margins/FCF with saved inputs are missing")
    require(any(cell["metric"] in {"revenue_yoy", "revenue_qoq"} and cell.get("formula") and
                (cell.get("value") is not None or cell.get("reason")) for cell in cells),
            "Missing deterministic revenue growth or explicit unavailable-baseline reason")
    return {"status": "passed", "numeric_ids": sorted(numeric_ids),
            "document_id": document["artifact_id"], "narrative_id": narrative["artifact_id"],
            "export_id": export["artifact_id"], "narrative_input_ids": sorted(inputs),
            "actual_periods": [{"start": start, "end": end} for start, end in periods],
            "unavailable_cells": sum(cell.get("value") is None for item in numeric
                                     for row in (item.get("data") or {}).get("rows", []) for cell in row.get("cells", [])),
            "latest_period_check": "alignment_only; latest-disclosed selection requires separate source audit"}


class SkillAcceptance(FilingAcceptance):
    def __init__(self, args):
        super().__init__(args)
        self.report.update({"scope": "browser_and_owned_result_chain", "session_name": SESSION_NAME,
                            "external_audit": {"skill_invocation_and_hash": "not_verified_by_this_runner",
                                               "narrative_quality": "requires_separate_review"}})
        for name in ("owned_chain", "workspace_skill", "table_source", "narrative_citation",
                     "document_citation_download", "formula_inputs", "download", "reload",
                     "owned_chain-after-reload", "workspace_skill-after-reload",
                     "table_source-after-reload", "narrative_citation-after-reload",
                     "document_citation_download-after-reload"):
            self.report["checks"][name] = {"status": "not_run"}

    def provision_session(self):
        sessions = self.api_get("/sessions/", {"agent_id": self.args.agent_id}).get("sessions", [])
        source = next((item for item in sessions if item["session"]["id"] == self.args.source_session_id), None)
        require(source is not None, "Source session was not found for this agent/user")
        config = model_reference(source)
        # Exactly one POST. A transport failure is not retried, even if the server
        # may already have committed a session whose ID we never received.
        response = self.api.post("/sessions/", json={"agent_id": self.args.agent_id,
                                 "name": SESSION_NAME, "chat_model_config": config})
        require(response.is_success, f"Test session creation failed (HTTP {response.status_code})")
        self.args.session_id = response.json()["session_id"]
        self.report["session_id"] = self.args.session_id
        self.report["model"] = {"model": config["model"], "type": config["type"],
                                "max_output_tokens_per_call": config["parameters"]["max_tokens"]}
        print(json.dumps({"created_session_id": self.args.session_id, **self.report["model"]}, ensure_ascii=False), flush=True)

    def run(self):
        try:
            require(self.args.run_real_model, "Requires explicit --run-real-model")
            require(len(self.args.prompts) == 1, "This acceptance authorizes exactly one browser prompt")
            self.provision_session()
        except Exception as error:
            self.report.update({"status": "failed", "failure": {"phase": "provision",
                                "type": type(error).__name__, "message": sanitized_error(error)}})
            self.api.close()
            path = self.output / "report.json"
            path.write_text(json.dumps(self.report, ensure_ascii=False, indent=2), encoding="utf-8")
            print(json.dumps({"status": "failed", "report": str(path)}, ensure_ascii=False))
            return 1
        return super().run()

    def test_sources(self, suffix=""):
        saved = self.api_get("/fin/results", {"session_id": self.args.session_id, "limit": "10"})["artifacts"]
        self.report["checks"]["owned_chain" + suffix] = validate_chain(saved)
        # Native Skill attachment is lazy on the first FinanceAgent reply. Check
        # AFTER that reply, using the browser's request context and authorized user.
        response = self.page.request.get(self.args.api_url + "/workspace/skill",
            params={"agent_id": self.args.agent_id, "session_id": self.args.session_id},
            headers={"X-User-ID": self.args.user_id})
        require(response.ok, f"Native workspace Skill GET failed (HTTP {response.status})")
        skills = response.json()
        require(isinstance(skills, list) and any(item.get("name") == "financial-analysis" for item in skills),
                "Default financial-analysis Skill is not attached to the test workspace")
        self.report["checks"]["workspace_skill" + suffix] = {
            "status": "passed", "name": "financial-analysis", "method": "native_authorized_browser_GET",
            "proves": "attachment_only", "maintenance": "not_exercised; no upload/remove/editor actions"}
        Acceptance.test_sources(self, suffix)
        with self.isolated_downloads("filing" + suffix):
            FilingAcceptance.test_sources(self, suffix)

    @contextmanager
    def isolated_downloads(self, name):
        # Every untrusted attachment gets a fresh directory, separate from the
        # trusted scripts and from the other attachment. Never execute its files.
        original = self.output
        directory = original / name
        directory.mkdir(exist_ok=False)
        self.output = directory
        try:
            yield
        finally:
            self.output = original

    def test_assistant_citations(self, suffix=""):
        Acceptance.test_assistant_citations(self, suffix)

    def test_formula(self):
        Acceptance.test_formula(self)
        require(self.report["checks"]["formula_inputs"]["status"] == "passed",
                "No clickable deterministic formula/input provenance was verified")

    def test_download(self):
        with self.isolated_downloads("excel"):
            Acceptance.test_download(self)

    def verify_workbook(self, path, export_id):
        result = Acceptance.verify_workbook(self, path, export_id)
        exported = self.api_get(f"/fin/results/{quote(export_id, safe='')}")
        inputs = [self.api_get(f"/fin/results/{quote(item, safe='')}") for item in exported["input_ids"]]
        documents = [item for item in inputs if item["kind"] == "document"]
        expected_cells = {(item["artifact_id"], cell["cell_id"]) for item in inputs
                          for row in (item.get("data") or {}).get("rows", [])
                          for cell in row.get("cells", []) if cell.get("value") is not None}
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            require("文档来源" in workbook.sheetnames and len(documents) == 1,
                    "Workbook lacks the original-document source sheet/input")
            iterator = workbook["文档来源"].iter_rows(values_only=True)
            header = next(iterator, ())
            rows = [dict(zip(header, row)) for row in iterator]
            require(len(rows) == 1, "Workbook must contain exactly one original-document source")
            document = documents[0]
            expected = {"文档结果ID": document["artifact_id"], "来源文档ID": document["document"]["source_document_id"],
                        "公司": "AAPL", "报告期末": document["document"]["period_end"],
                        "Accession": document["document"]["accession"], "来源URL": document["document"]["source_url"],
                        "SHA-256": document["document"]["sha256"], "字节数": document["document"]["byte_size"]}
            require(all(rows[0].get(key) == value for key, value in expected.items()),
                    "Workbook document source metadata differs from the pinned owned artifact")
            observed = []
            narrative_ids = set()
            for sheet in workbook:
                iterator = sheet.iter_rows(values_only=True)
                header = next(iterator, ())
                if {"结果ID", "单元格ID", "精确值"}.issubset(header):
                    positions = {label: index for index, label in enumerate(header)}
                    observed.extend((row[positions["结果ID"]], row[positions["单元格ID"]])
                                    for row in iterator if row[positions["精确值"]] is not None)
                elif header == ("项目", "内容"):
                    narrative_ids.update(row[1] for row in iterator if row[0] == "结果ID")
            require(set(observed) == expected_cells and len(observed) == len(expected_cells),
                    "Workbook omits or duplicates stored non-null financial cells")
            require({item["artifact_id"] for item in inputs if item["kind"] == "narrative"}.issubset(narrative_ids),
                    "Workbook omits the saved analysis sheet")
            return {**result, "all_non_null_cells_exported": True, "document_source": expected,
                    "analysis_ids": sorted(narrative_ids)}
        finally:
            workbook.close()


def arguments(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-real-model", action="store_true")
    parser.add_argument("--agent-id", required=True)
    parser.add_argument("--source-session-id", required=True)
    parser.add_argument("--api-url", type=base_url, default="http://127.0.0.1:8000")
    parser.add_argument("--web-url", type=base_url, default="http://127.0.0.1:8080")
    parser.add_argument("--user-id", default="dev")
    parser.add_argument("--output-dir", type=Path, default=Path("/e2e-output"))
    args = parser.parse_args(argv)
    if not args.run_real_model:
        parser.error("Requires --run-real-model; creates one test session and spends provider tokens")
    if not all(value.strip() for value in (args.agent_id, args.source_session_id, args.user_id)):
        parser.error("Agent, source session and user IDs must be nonempty")
    args.session_id, args.multi_turn, args.headed = "", True, False
    args.ui_timeout, args.turn_timeout, args.prompts = 45, 600, [PROMPT]
    return args


def main(argv=None):
    return SkillAcceptance(arguments(argv)).run()


if __name__ == "__main__":
    raise SystemExit(main())
