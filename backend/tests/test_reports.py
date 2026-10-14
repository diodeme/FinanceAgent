"""Offline workbook tests: synthetic immutable artifacts, temporary files only.

Run: python -m unittest tests.test_reports -v (from backend with dependencies).
No PostgreSQL, HTTP, provider, or model calls are made.
"""

from datetime import date, datetime, timezone
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from finagent.contracts import (
    Artifact, Cell, Citation, DocumentCitation, Evidence, ExportRequest, FinanceError, FinancialRow, TableData,
)
from finagent.reports import excel_number, generate_report, safe_text, write_workbook
from finagent.sec_client import CONTACT_WARNING
from finagent.warnings import CROSS_COMPANY_WARNING, CROSS_COMPANY_QUARTERS_ALIAS, ORIGINAL_DOCUMENT_WARNING


NOW = datetime(2025, 5, 2, tzinfo=timezone.utc)
EXACT = "1234567890123456789.1234"


def synthetic_artifacts():
    evidence = Evidence(
        source_id="test-source", fact_id=101, company="AAPL", concept="us-gaap:Revenues",
        metric="revenue", period_start=date(2025, 1, 1), period_end=date(2025, 3, 31),
        value="123456.7890", unit="USD", accession="0000320193-25-999999", form="10-Q",
        filed_at=date(2025, 5, 1), retrieved_at=NOW,
        api_url="https://data.sec.gov/api/xbrl/companyfacts/CIK0000320193.json",
        filing_url="https://www.sec.gov/Archives/edgar/data/320193/synthetic-test-only.htm",
    )
    precise = evidence.model_copy(update={"source_id": "test-precise-source", "fact_id": 102,
                                          "metric": "total_assets", "concept": "us-gaap:Assets", "value": EXACT})
    row = FinancialRow(
        company="AAPL", period_start=evidence.period_start, period_end=evidence.period_end,
        period_label="FY2025 Q2", cells=[
            Cell(cell_id="test-revenue-cell", metric="revenue", label="Revenue", value=evidence.value,
                 unit="USD", source_ids=[evidence.source_id], formula="=SUM(A1:A2)"),
            Cell(cell_id="test-precise-cell", metric="total_assets", label="Total assets", value=EXACT,
                 unit="USD", source_ids=[precise.source_id]),
            Cell(cell_id="test-missing-cell", metric="net_income", label="Net income", value=None,
                 unit="USD", status="missing", reason="No synthetic source available"),
            Cell(cell_id="test-unrequested-cell", metric="cost_of_revenue", label="Do not display",
                 value="999", unit="USD"),
        ],
    )
    history = row.model_copy(deep=True)
    history.period_label = "HISTORY ONLY"
    history.period_end = date(2024, 3, 31)
    history.cells = [Cell(cell_id="test-history-cell", metric="revenue", label="Historical value",
                          value="887766", unit="USD")]
    dataset = Artifact(
        artifact_id="test-dataset", run_id="test-data-run", kind="dataset", title="Synthetic results",
        created_at=NOW, warnings=["Shared warning"],
        data=TableData(rows=[row], history_rows=[history], sources={evidence.source_id: evidence, precise.source_id: precise},
                       specification={"requested_metrics": ["revenue", "total_assets", "net_income"]},
                       warnings=["Shared warning", "Synthetic evidence only"]),
    )
    narrative = Artifact(
        artifact_id="test-narrative", run_id="test-analysis-run", kind="narrative", title="Optional analysis",
        created_at=NOW, text="Synthetic analysis\n=HYPERLINK(\"https://example.invalid\",\"not executable\")",
        input_ids=[dataset.artifact_id],
        citations=[Citation(artifact_id=dataset.artifact_id, cell_id="test-revenue-cell", label="Revenue evidence")],
    )
    return dataset, narrative


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="finagent-workbook-test-")
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "nested" / "report.xlsx"

    def open_report(self, dataset, analyses=None, title="Synthetic report"):
        write_workbook(self.path, title, [dataset], analyses or [])
        self.assertTrue(self.path.is_file())
        self.assertFalse(self.path.with_suffix(".tmp").exists())
        workbook = load_workbook(self.path, data_only=False)
        self.addCleanup(workbook.close)
        return workbook

    def test_safe_text_blocks_formula_prefixes_and_illegal_xml_controls(self):
        for value in ("=1+1", "+SUM(A1)", "-1+2", "@SUM(A1)", " \t=1+1", "\n+1"):
            with self.subTest(value=value):
                self.assertEqual(safe_text(value), "'" + value)
        self.assertEqual(safe_text(None), "")
        self.assertEqual(safe_text("safe\x00\x01\x0b\x1f text"), "safe text")
        self.assertEqual(safe_text("plain\ttext\nline"), "plain\ttext\nline")
        self.assertEqual(len(safe_text("x" * 40000)), 32767)

    def test_excel_numbers_preserve_high_precision_and_reject_nonfinite_values(self):
        self.assertIsNone(excel_number(None))
        self.assertEqual(excel_number("123456.7890"), 123456.789)
        self.assertEqual(excel_number("-12.50"), -12.5)
        self.assertEqual(excel_number(EXACT), EXACT)
        for value in ("not a number", "NaN", "Infinity", "-Infinity"):
            with self.subTest(value=value), self.assertRaises(FinanceError) as caught:
                excel_number(value)
            self.assertEqual(caught.exception.code, "invalid_number")

    def test_workbook_exact_values_cell_sources_and_no_visible_history(self):
        dataset, _ = synthetic_artifacts()
        original = dataset.model_dump(mode="json")
        workbook = self.open_report(dataset)
        self.assertEqual(workbook.sheetnames, ["说明", "来源", "覆盖与限制", "数据1"])
        sheet = workbook["数据1"]
        self.assertEqual(sheet.max_row, 4)  # header plus three requested cells
        self.assertEqual(sheet["G2"].value, 123456.789)
        self.assertEqual(sheet["G2"].data_type, "n")
        self.assertEqual(sheet["H2"].value, "123456.7890")
        self.assertEqual(sheet["H2"].data_type, "s")
        self.assertEqual(sheet["G3"].value, EXACT)
        self.assertEqual(sheet["G3"].data_type, "s")
        self.assertEqual(sheet["H3"].value, EXACT)
        self.assertIsNone(sheet["G4"].value)
        self.assertEqual(sheet["J4"].value, "missing")
        self.assertEqual(sheet["K2"].value, "'=SUM(A1:A2)")
        self.assertEqual(sheet["L2"].value, "test-source")
        self.assertEqual(sheet["M2"].value, "test-revenue-cell")
        self.assertEqual(sheet["L3"].value, "test-precise-source")
        self.assertEqual(sheet["M3"].value, "test-precise-cell")
        sources = {row[1]: row for row in workbook["来源"].iter_rows(min_row=2, values_only=True)}
        source = dataset.data.sources["test-source"]
        source_row = sources[source.source_id]
        self.assertEqual(source_row[0], dataset.artifact_id)
        self.assertEqual(source_row[2], source.fact_id)
        self.assertEqual(source_row[7], source.value)
        self.assertEqual(source_row[11:14], (source.api_url, source.filing_url, "metadata_only"))
        self.assertEqual(sources["test-precise-source"][7], EXACT)
        for name in ("数据1", "覆盖与限制"):
            visible = [cell.value for row in workbook[name] for cell in row]
            self.assertNotIn("HISTORY ONLY", visible)
            self.assertNotIn("test-history-cell", visible)
            self.assertNotIn("test-unrequested-cell", visible)
            self.assertNotIn("Do not display", visible)
        coverage = list(workbook["覆盖与限制"].values)
        self.assertTrue(any(row[4:6] == ("missing", "No synthetic source available") for row in coverage))
        self.assertEqual(sum(row[-1] == "Shared warning" for row in coverage), 1)
        self.assertEqual(dataset.model_dump(mode="json"), original)
        self.assertEqual(sheet.freeze_panes, "A2")

    def test_optional_narrative_retains_citations_and_never_creates_formula_cells(self):
        dataset, narrative = synthetic_artifacts()
        dataset.title = "=1+1"
        dataset.data.rows[0].period_label = "+period"
        dataset.data.rows[0].cells[0].label = "@metric"
        dataset.data.rows[0].cells[2].reason = "\t=reason"
        dataset.data.sources["test-source"].concept = "=concept"
        dataset.warnings.append("-warning")
        narrative.title = "+analysis"
        narrative.citations[0].label = "=citation label"
        workbook = self.open_report(dataset, [narrative], title="=report\x00")
        self.assertIn("分析1", workbook.sheetnames)
        self.assertEqual(workbook["说明"]["B2"].value, "'=report")
        self.assertEqual(workbook["分析1"]["B3"].value, "'+analysis")
        analysis_values = [row[1] for row in workbook["分析1"].values]
        self.assertIn("test-dataset / test-revenue-cell / =citation label", analysis_values)
        self.assertIn("Synthetic analysis", analysis_values)
        self.assertTrue(any(str(value).startswith("'=HYPERLINK") for value in analysis_values))
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f", f"Executable formula in {sheet.title}!{cell.coordinate}")

    def test_workbook_warning_aliases_scope_and_collection_diagnostics(self):
        dataset, narrative = synthetic_artifacts()
        dataset.warnings += [CROSS_COMPANY_QUARTERS_ALIAS, CONTACT_WARNING, ORIGINAL_DOCUMENT_WARNING]
        dataset.data.warnings += [CROSS_COMPANY_WARNING, "Capex scope limitation", "Missing metric", "Conflicting values"]
        narrative.warnings = [CONTACT_WARNING, CROSS_COMPANY_WARNING, "Narrative-only limitation"]
        before = [a.model_dump(mode="json") for a in [dataset, narrative]]
        workbook = self.open_report(dataset, [narrative])
        warnings = [row[-1] for row in workbook["覆盖与限制"].values if row[4] == "warning"]
        self.assertNotIn(CROSS_COMPANY_WARNING, warnings)
        self.assertNotIn(CROSS_COMPANY_QUARTERS_ALIAS, warnings)
        self.assertNotIn(CONTACT_WARNING, warnings)
        for expected in [ORIGINAL_DOCUMENT_WARNING, "Capex scope limitation", "Missing metric", "Conflicting values", "Narrative-only limitation", "Synthetic evidence only"]:
            self.assertIn(expected, warnings)
        self.assertEqual([a.model_dump(mode="json") for a in [dataset, narrative]], before)

    def test_company_and_source_identifiers_are_also_safe_text(self):
        dataset, _ = synthetic_artifacts()
        dataset.data.rows[0].company = "=company"
        dataset.data.sources["test-source"].company = "=source_company"
        dataset.data.sources["test-source"].source_id = "=source_id"
        workbook = self.open_report(dataset)
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f", f"Executable formula in {sheet.title}!{cell.coordinate}")


class ReportSelectionTests(unittest.IsolatedAsyncioTestCase):
    async def test_narrative_only_export_includes_its_owned_evidence(self):
        dataset, narrative = synthetic_artifacts()
        fixtures = {item.artifact_id: item for item in (dataset, narrative)}
        calls = []

        async def get_owned(user_id, artifact_id):
            calls.append((user_id, artifact_id))
            return fixtures[artifact_id]

        exported = Artifact(artifact_id="test-export", run_id="test-export-run", kind="export",
                            title="Test export", created_at=NOW)
        with tempfile.TemporaryDirectory(prefix="finagent-export-test-") as directory:
            root = Path(directory)
            with patch("finagent.reports.get_artifact", side_effect=get_owned), \
                 patch("finagent.reports.report_root", return_value=root), \
                 patch("finagent.reports.save_artifact", new=AsyncMock(return_value=exported)) as save:
                result = await generate_report("test-owner", "test-export-run", ExportRequest(
                    analysis_ids=[narrative.artifact_id], title="Test export",
                ))
                self.assertEqual(result, exported)
                self.assertEqual(calls, [("test-owner", narrative.artifact_id), ("test-owner", dataset.artifact_id)])
                self.assertEqual(save.await_count, 1)
                args, kwargs = save.await_args
                self.assertEqual(args[:3], ("test-owner", "test-export-run", "export"))
                self.assertEqual(kwargs["input_ids"], [dataset.artifact_id, narrative.artifact_id])
                self.assertIn("Synthetic evidence only", kwargs["warnings"])
                self.assertNotIn(CROSS_COMPANY_WARNING, kwargs["warnings"])
                path = Path(kwargs["path"])
                self.assertTrue(path.is_file())
                self.assertEqual(path.parent, root)
                workbook = load_workbook(path)
                try:
                    self.assertIn("数据1", workbook.sheetnames)
                    self.assertIn("分析1", workbook.sheetnames)
                finally:
                    workbook.close()

    async def test_export_company_scope_includes_narrative_dependencies(self):
        dataset, narrative = synthetic_artifacts()
        other = dataset.model_copy(deep=True)
        other.artifact_id = "other-company"
        other.data.rows[0].company = "MSFT"
        other.data.warnings += ["Dependency table-only limitation", CONTACT_WARNING]
        dataset.warnings += [CROSS_COMPANY_QUARTERS_ALIAS, ORIGINAL_DOCUMENT_WARNING, CONTACT_WARNING]
        narrative.input_ids = [other.artifact_id]
        narrative.citations = [Citation(artifact_id=other.artifact_id, cell_id="test-revenue-cell", label="Other evidence")]
        narrative.warnings = [CROSS_COMPANY_WARNING, CONTACT_WARNING]
        fixtures = {a.artifact_id: a for a in [dataset, other, narrative]}
        before = {key: a.model_dump(mode="json") for key, a in fixtures.items()}
        async def get_owned(user_id, artifact_id):
            self.assertEqual(user_id, "test-owner")
            return fixtures[artifact_id]
        with tempfile.TemporaryDirectory(prefix="finagent-export-warning-test-") as directory:
            with patch("finagent.reports.get_artifact", side_effect=get_owned), \
                 patch("finagent.reports.report_root", return_value=Path(directory)), \
                 patch("finagent.reports.save_artifact", new=AsyncMock()) as save:
                await generate_report("test-owner", "run", ExportRequest(data_ids=[dataset.artifact_id], analysis_ids=[narrative.artifact_id]))
                warnings = save.await_args.kwargs["warnings"]
                self.assertEqual(warnings.count(CROSS_COMPANY_WARNING), 1)
                self.assertNotIn(CROSS_COMPANY_QUARTERS_ALIAS, warnings)
                self.assertNotIn(CONTACT_WARNING, warnings)
                self.assertIn("Dependency table-only limitation", warnings)
                self.assertIn(ORIGINAL_DOCUMENT_WARNING, warnings)
                workbook = load_workbook(save.await_args.kwargs["path"])
                try:
                    sheet_warnings = [row[-1] for row in workbook["覆盖与限制"].values if row[4] == "warning"]
                    self.assertEqual(sheet_warnings, warnings)
                finally:
                    workbook.close()
        self.assertEqual({key: a.model_dump(mode="json") for key, a in fixtures.items()}, before)

    async def test_document_only_export_preserves_hash_source_and_safe_citations(self):
        from tests.test_filings import artifact_fixture
        document = artifact_fixture()
        narrative = Artifact(
            artifact_id="document-analysis", run_id="analysis-run", kind="narrative", title="Document-only",
            created_at=NOW, input_ids=[document.artifact_id], text="=unsafe formula\nDocument-based observation",
            citations=[DocumentCitation(document_id=document.artifact_id, label="=source")],
        )
        fixtures = {item.artifact_id: item for item in [document, narrative]}
        calls = []
        async def owned(user, artifact_id):
            calls.append((user, artifact_id))
            return fixtures[artifact_id]
        with tempfile.TemporaryDirectory() as directory:
            with patch("finagent.reports.get_artifact", side_effect=owned), patch("finagent.reports.report_root", return_value=Path(directory)), patch("finagent.reports.save_artifact", new=AsyncMock()) as save:
                await generate_report("owner", "run", ExportRequest(analysis_ids=[narrative.artifact_id]))
                self.assertEqual(calls, [("owner", narrative.artifact_id), ("owner", document.artifact_id)])
                self.assertEqual(save.await_args.kwargs["input_ids"], [document.artifact_id, narrative.artifact_id])
                workbook = load_workbook(save.await_args.kwargs["path"])
                try:
                    self.assertIn("文档来源", workbook.sheetnames)
                    self.assertNotIn("数据1", workbook.sheetnames)
                    values = [cell.value for sheet in workbook for row in sheet for cell in row]
                    self.assertIn(document.document.sha256, values)
                    self.assertIn(document.document.source_url, values)
                    self.assertIn("document_fixture / =source", values)
                    self.assertFalse(any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row))
                finally:
                    workbook.close()
        with patch("finagent.reports.get_artifact", new=AsyncMock(side_effect=[narrative, FinanceError("not_found", "Foreign document")])), patch("finagent.reports.write_workbook") as write:
            with self.assertRaises(FinanceError):
                await generate_report("other", "run", ExportRequest(analysis_ids=[narrative.artifact_id]))
            write.assert_not_called()

    async def test_document_scope_qualifies_old_warning_without_mutating_snapshot(self):
        from tests.test_filings import artifact_fixture
        from finagent.warnings import artifact_companies, artifact_warnings
        dataset, narrative = synthetic_artifacts()
        document = artifact_fixture()
        document.document.company = "MSFT"
        dataset.warnings += [ORIGINAL_DOCUMENT_WARNING, "Unknown limitation stays verbatim"]
        before = dataset.model_dump(mode="json")
        warnings = artifact_warnings([dataset, narrative, document])
        self.assertEqual(artifact_companies([dataset, document]), {"AAPL", "MSFT"})
        self.assertIn(CROSS_COMPANY_WARNING, warnings)
        self.assertIn("Unknown limitation stays verbatim", warnings)
        self.assertNotIn(ORIGINAL_DOCUMENT_WARNING, warnings)
        self.assertTrue(any("metadata-only provenance" in warning for warning in warnings))
        self.assertEqual(dataset.model_dump(mode="json"), before)

    async def test_wrong_kind_or_foreign_input_cannot_publish_export(self):
        _, narrative = synthetic_artifacts()
        for error in (None, FinanceError("not_found", "Not owned")):
            get = AsyncMock(return_value=narrative, side_effect=error)
            with self.subTest(foreign=error is not None), \
                 patch("finagent.reports.get_artifact", new=get), \
                 patch("finagent.reports.write_workbook") as write, \
                 patch("finagent.reports.save_artifact", new=AsyncMock()) as save:
                with self.assertRaises(FinanceError) as caught:
                    await generate_report("test-owner", "test-export-run", ExportRequest(data_ids=["test-input"]))
                self.assertEqual(caught.exception.code, "not_found" if error else "invalid_result")
                get.assert_awaited_once_with("test-owner", "test-input")
                write.assert_not_called()
                save.assert_not_awaited()


if __name__ == "__main__":
    unittest.main()
