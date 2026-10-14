"""Offline acceptance-runner regression tests. Never start a browser/API/model.

Run in the existing finance-e2e image with --network none and the e2e directory
mounted read-only; temporary generated workbooks stay in the container's /tmp.
"""
import ast
from contextlib import redirect_stderr
from copy import deepcopy
import io
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook

import filing_acceptance as filing_runner
import skill_acceptance as runner
from acceptance import Acceptance
from filing_acceptance import FilingAcceptance


def artifacts():
    periods = [("2025-04-01", "2025-06-30"), ("2025-07-01", "2025-09-30"),
               ("2025-10-01", "2025-12-31"), ("2026-01-01", "2026-03-31")]
    rows = [{"company": "AAPL", "period_start": start, "period_end": end,
             "cells": [{"cell_id": f"base-{index}-{metric}", "metric": metric, "value": "12.50"}
                       for metric in sorted(runner.BASE_METRICS)]}
            for index, (start, end) in enumerate(periods)]
    calculated = [{"cell_id": metric, "metric": metric, "value": "0.25", "formula": "a / b",
                   "inputs": [{"value": "1"}, {"value": "4"}]}
                  for metric in sorted(runner.DERIVED_METRICS | {"revenue_yoy"})]
    return [
        {"artifact_id": "dataset-1", "kind": "dataset", "data": {"rows": rows}},
        {"artifact_id": "table-1", "kind": "table", "data": {
            "specification": {"operation": "calculate"}, "rows": [{**rows[-1], "cells": calculated}]}},
        {"artifact_id": "document-1", "kind": "document", "document": {
            "company": "AAPL", "source_document_id": "source-1", "period_end": "2026-03-31",
            "sha256": "a" * 64, "byte_size": 12, "accession": "accession-1",
            "source_url": "https://www.sec.gov/example.html"}},
        {"artifact_id": "narrative-1", "kind": "narrative", "text": "Offline fixture; not a live report.",
         "input_ids": ["dataset-1", "table-1", "document-1"],
         "citations": [{"artifact_id": "table-1", "cell_id": "net_margin"}, {"document_id": "document-1"}]},
        {"artifact_id": "export-1", "kind": "export", "filename": "report.xlsx",
         "input_ids": ["dataset-1", "table-1", "document-1", "narrative-1"]},
    ]


def source_session():
    return {"is_running": False, "session": {"id": "existing", "config": {"chat_model_config": {
        "type": "existing-provider", "model": "existing-model", "credential_id": "reference-only",
        "parameters": {"temperature": 0.3, "max_tokens": 9000}, "ignored": "not-copied"}}}}


class RunnerTests(unittest.TestCase):
    def test_ast_all_e2e_python(self):
        for path in Path(__file__).parent.glob("*.py"):
            with self.subTest(file=path.name):
                ast.parse(path.read_text(encoding="utf-8"), filename=str(path))

    def test_opt_in_guard_before_harness_or_session_creation(self):
        with patch.object(runner, "SkillAcceptance") as harness, redirect_stderr(io.StringIO()):
            with self.assertRaises(SystemExit) as error:
                runner.main(["--agent-id", "agent", "--source-session-id", "source"])
        self.assertEqual(error.exception.code, 2)
        harness.assert_not_called()

    def test_single_prompt_limits(self):
        args = runner.arguments(["--run-real-model", "--agent-id", "agent", "--source-session-id", "source"])
        self.assertEqual(args.prompts, [runner.PROMPT])
        self.assertEqual(args.turn_timeout, 600)
        self.assertEqual(len(runner.BASE_METRICS), 10)
        self.assertLessEqual(len(runner.BASE_METRICS), 12)
        self.assertIn("600至900字", runner.PROMPT)
        self.assertIn("financial-analysis", runner.PROMPT)
        self.assertEqual(runner.SESSION_NAME, "Skill 综合财务分析实测")

    def test_model_reference_preserves_parameters_without_mutating_source(self):
        self.assertIs(runner.model_reference, filing_runner.model_reference)
        for limit in (9000, 1200, None):
            with self.subTest(max_tokens=limit):
                source = source_session()
                parameters = source["session"]["config"]["chat_model_config"]["parameters"]
                parameters.update({"max_tokens": limit, "nested": {"stop": ["END"]}})
                original = deepcopy(source)
                config = runner.model_reference(source)
                self.assertEqual(source, original)
                self.assertEqual(set(config), {"type", "model", "credential_id", "parameters"})
                self.assertEqual(config["model"], "existing-model")
                self.assertEqual(config["parameters"], parameters)
                config["parameters"]["nested"]["stop"].append("OTHER")
                config["parameters"]["max_tokens"] = 1
                self.assertEqual(source, original)
        source["is_running"] = True
        with self.assertRaises(AssertionError):
            runner.model_reference(source)

    def test_model_reference_does_not_inject_missing_output_limit(self):
        for parameters in ({}, {"temperature": 0.3}):
            source = source_session()
            source["session"]["config"]["chat_model_config"]["parameters"] = parameters
            original = deepcopy(source)
            config = runner.model_reference(source)
            self.assertEqual(config["parameters"], parameters)
            self.assertNotIn("max_tokens", config["parameters"])
            self.assertEqual(source, original)
        del source["session"]["config"]["chat_model_config"]["parameters"]
        self.assertEqual(runner.model_reference(source)["parameters"], {})

    def test_filing_opt_in_guard_before_api_or_browser(self):
        with patch("sys.argv", ["filing_acceptance.py", "--agent-id", "agent",
                                "--source-session-id", "source"]), \
                patch.object(filing_runner.httpx, "Client") as client, \
                patch.object(filing_runner, "FilingAcceptance") as harness, \
                redirect_stderr(io.StringIO()):
            with self.assertRaises(SystemExit) as error:
                filing_runner.main()
        self.assertEqual(error.exception.code, 2)
        client.assert_not_called()
        harness.assert_not_called()

    def test_filing_session_creation_preserves_config_without_source_writes(self):
        for parameters in ({}, {"temperature": 0.3, "max_tokens": 9000}):
            source = source_session()
            source["session"]["config"]["chat_model_config"]["parameters"] = parameters
            original = deepcopy(source)
            with patch("sys.argv", ["filing_acceptance.py", "--run-real-model",
                                    "--agent-id", "agent", "--source-session-id", "existing"]), \
                    patch.object(filing_runner.httpx, "Client") as client_factory, \
                    patch.object(filing_runner, "FilingAcceptance") as harness, \
                    patch("builtins.print"):
                client = client_factory.return_value.__enter__.return_value
                client.get.return_value.json.return_value = {"sessions": [source]}
                client.post.return_value.json.return_value = {"session_id": "new-test"}
                harness.return_value.run.return_value = 0
                self.assertEqual(filing_runner.main(), 0)
            client.post.assert_called_once_with("/sessions/", json={
                "agent_id": "agent", "name": "SEC 原文分析实测",
                "chat_model_config": runner.model_reference(source)})
            self.assertEqual([call[0] for call in client.method_calls], ["get", "post"])
            self.assertEqual(source, original)

    def test_skill_session_creation_preserves_config_without_source_writes(self):
        for parameters in ({}, {"temperature": 0.3, "max_tokens": 9000}):
            source = source_session()
            source["session"]["config"]["chat_model_config"]["parameters"] = parameters
            original = deepcopy(source)
            harness = object.__new__(runner.SkillAcceptance)
            harness.args = SimpleNamespace(agent_id="agent", source_session_id="existing", session_id="")
            harness.report = {}
            harness.api_get = Mock(return_value={"sessions": [source]})
            harness.api = Mock()
            harness.api.post.return_value.is_success = True
            harness.api.post.return_value.json.return_value = {"session_id": "new-test"}
            with patch("builtins.print"):
                harness.provision_session()
            harness.api.post.assert_called_once_with("/sessions/", json={
                "agent_id": "agent", "name": runner.SESSION_NAME,
                "chat_model_config": runner.model_reference(source)})
            self.assertEqual([call[0] for call in harness.api.method_calls], ["post"])
            self.assertEqual(harness.report["model"]["max_output_tokens_per_call"], parameters.get("max_tokens"))
            self.assertEqual(source, original)

    def test_inline_secret_is_rejected(self):
        source = source_session()
        source["session"]["config"]["chat_model_config"]["parameters"]["nested"] = {"api_key": "placeholder"}
        with self.assertRaisesRegex(AssertionError, "inline credentials"):
            runner.model_reference(source)

    def test_valid_owned_chain_and_explicit_missing_growth(self):
        saved = artifacts()
        check = runner.validate_chain(saved)
        self.assertEqual(check["status"], "passed")
        self.assertEqual(len(check["actual_periods"]), 4)
        growth = next(cell for cell in saved[1]["data"]["rows"][0]["cells"] if cell["metric"] == "revenue_yoy")
        growth.update({"value": None, "reason": "No comparable prior quarter"})
        self.assertEqual(runner.validate_chain(saved)["status"], "passed")

    def test_missing_chain_parts_never_pass(self):
        for kind in ("dataset", "table", "document", "narrative", "export"):
            with self.subTest(kind=kind), self.assertRaises(AssertionError):
                runner.validate_chain([item for item in artifacts() if item["kind"] != kind])
        for inputs in (["dataset-1", "table-1"], ["document-1"]):
            saved = artifacts()
            saved[3]["input_ids"] = inputs
            with self.subTest(inputs=inputs), self.assertRaises(AssertionError):
                runner.validate_chain(saved)

    def test_extra_document_misalignment_missing_calculation_or_export_rejected(self):
        for mutation in (
            lambda saved: saved.append({**saved[2], "artifact_id": "extra-doc"}),
            lambda saved: saved[2]["document"].update({"period_end": "2025-12-31"}),
            lambda saved: saved[1]["data"].update({"specification": {}}),
            lambda saved: saved[4].update({"input_ids": ["dataset-1", "document-1"]}),
            lambda saved: saved[3].update({"citations": [{"document_id": "document-1"}]}),
        ):
            saved = artifacts()
            mutation(saved)
            with self.assertRaises(AssertionError):
                runner.validate_chain(saved)

    def test_session_post_once_without_retry(self):
        harness = object.__new__(runner.SkillAcceptance)
        harness.args = SimpleNamespace(agent_id="agent", source_session_id="existing", session_id="")
        harness.report = {}
        harness.api_get = Mock(return_value={"sessions": [source_session()]})
        harness.api = Mock()
        harness.api.post.side_effect = RuntimeError("offline simulated transport failure")
        with self.assertRaises(RuntimeError):
            harness.provision_session()
        self.assertEqual(harness.api.post.call_count, 1)
        self.assertEqual(harness.api.post.call_args.kwargs["json"]["name"], runner.SESSION_NAME)

    def test_reuses_harness_methods(self):
        harness = object.__new__(runner.SkillAcceptance)
        harness.report = {"checks": {"formula_inputs": {"status": "passed"}}}
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        harness.output = Path(directory.name)
        for name in ("test_download", "test_assistant_citations", "test_formula"):
            with patch.object(Acceptance, name) as inherited:
                getattr(harness, name)()
                inherited.assert_called_once()
        self.assertIs(runner.SkillAcceptance.__mro__[1], FilingAcceptance)

    def test_native_attachment_and_both_citation_harnesses_before_after_reload(self):
        harness = object.__new__(runner.SkillAcceptance)
        harness.args = SimpleNamespace(api_url="http://localhost:8000", agent_id="agent", session_id="test", user_id="owner")
        harness.report = {"checks": {}}
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        harness.output = Path(directory.name)
        harness.api_get = Mock(return_value={"artifacts": artifacts()})
        harness.page = Mock()
        response = harness.page.request.get.return_value
        response.ok = True
        response.json.return_value = [{"name": "financial-analysis"}]
        with patch.object(Acceptance, "test_sources") as numeric, patch.object(FilingAcceptance, "test_sources") as filing:
            for suffix in ("", "-after-reload"):
                harness.test_sources(suffix)
                numeric.assert_called_with(harness, suffix)
                filing.assert_called_with(harness, suffix)
                self.assertEqual(harness.report["checks"]["workspace_skill" + suffix]["proves"], "attachment_only")
            self.assertEqual(numeric.call_count, 2)
            self.assertEqual(filing.call_count, 2)
            response.json.return_value = []
            with self.assertRaisesRegex(AssertionError, "not attached"):
                harness.test_sources()
            self.assertEqual(numeric.call_count, 2)
        harness.page.request.get.assert_called_with("http://localhost:8000/workspace/skill",
            params={"agent_id": "agent", "session_id": "test"}, headers={"X-User-ID": "owner"})
        harness.page.request.post.assert_not_called()
        harness.page.request.delete.assert_not_called()

    def test_workbook_exact_values_completeness_document_metadata(self):
        saved = artifacts()
        harness = object.__new__(runner.SkillAcceptance)
        by_id = {item["artifact_id"]: item for item in saved}
        harness.api_get = lambda path: by_id[path.rsplit("/", 1)[-1]]
        book = Workbook()
        book.active.title = "来源"
        book.create_sheet("覆盖与限制")
        sheet = book.create_sheet("数据1")
        sheet.append(["结果ID", "单元格ID", "值", "精确值"])
        for item in saved[:2]:
            for row in item["data"]["rows"]:
                for cell in row["cells"]:
                    sheet.append([item["artifact_id"], cell["cell_id"], cell["value"], cell["value"]])
        docs = book.create_sheet("文档来源")
        docs.append(["文档结果ID", "来源文档ID", "公司", "报告期末", "Accession", "来源URL", "SHA-256", "字节数"])
        doc = saved[2]["document"]
        docs.append(["document-1", "source-1", "AAPL", doc["period_end"], doc["accession"], doc["source_url"], doc["sha256"], 12])
        analysis = book.create_sheet("分析1")
        analysis.append(["项目", "内容"])
        analysis.append(["结果ID", "narrative-1"])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "offline.xlsx"
            book.save(path)
            result = harness.verify_workbook(path, "export-1")
            self.assertTrue(result["all_non_null_cells_exported"])
            self.assertEqual(result["checked_values"], 45)
            docs["G2"] = "wrong-hash"
            book.save(path)
            with self.assertRaisesRegex(AssertionError, "metadata differs"):
                harness.verify_workbook(path, "export-1")
            docs["G2"] = doc["sha256"]
            sheet["D2"] = "99.00"
            book.save(path)
            with self.assertRaisesRegex(AssertionError, "exact-value column differs"):
                harness.verify_workbook(path, "export-1")
            sheet["D2"] = "12.50"
            sheet.delete_rows(2)
            book.save(path)
            with self.assertRaisesRegex(AssertionError, "omits or duplicates"):
                harness.verify_workbook(path, "export-1")
        book.close()


if __name__ == "__main__":
    unittest.main()
